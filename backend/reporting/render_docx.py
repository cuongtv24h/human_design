"""Word (.docx) from ``ReportDocument`` (plan P0-6) — python-docx.

Same content as the PDF/Markdown: cover with subject info + chart facts,
BodyGraph image, then every included section. Built from Word's own styles
(Title, Heading 1–3, List Bullet…) so coaches can restyle or edit freely.
"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .contract import ReportDocument
from .mdblocks import Code, Heading, Inline, ListBlock, Paragraph as MdParagraph, Quote, Rule, Table as MdTable
from .mdblocks import parse_blocks
from .render_common import (
    Theme, bodygraph_png, breaks_between_sections, chart_facts, included_sections, mode_label, prepared_on,
    subject_lines,
)

BODY_FONT = "Arial"  # ships with Word/Google Docs/LibreOffice and covers Vietnamese


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _set_font(style, size: float, color: str | None = None, bold: bool | None = None) -> None:
    style.font.name = BODY_FONT
    style.font.size = Pt(size)
    if color:
        style.font.color.rgb = _rgb(color)
    if bold is not None:
        style.font.bold = bold
    # East-Asian / complex-script slots too, otherwise Word may substitute fonts.
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), BODY_FONT)


def _setup_styles(doc, theme: Theme) -> None:
    styles = doc.styles
    _set_font(styles["Normal"], 11, theme.ink)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.25
    _set_font(styles["Title"], 24, theme.primary_dark, True)
    for name, size, color in (("Heading 1", 17, theme.primary_dark), ("Heading 2", 13.5, theme.primary),
                              ("Heading 3", 12, theme.ink)):
        _set_font(styles[name], size, color, True)
        styles[name].paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        styles[name].paragraph_format.space_after = Pt(6)
        styles[name].paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Bullet 2", "List Bullet 3"):
        _set_font(styles[name], 11, theme.ink)
        styles[name].paragraph_format.space_after = Pt(3)
    _set_font(styles["Quote"], 11, theme.ink)
    styles["Quote"].font.italic = False


def _add_runs(paragraph, inline: Inline) -> None:
    for span in inline:
        run = paragraph.add_run(span.text)
        run.bold = span.bold or None
        run.italic = span.italic or None
        if span.code:
            run.font.name = "Consolas"


def _add_lines(paragraph, lines: list[Inline]) -> None:
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break(WD_BREAK.LINE)
        _add_runs(paragraph, line)


# OOXML requires a fixed child order inside w:pPr / w:tcPr (Word rejects the file otherwise).
_PPR_AFTER_SHD = ("w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                  "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
                  "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                  "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                  "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
_PPR_AFTER_PBDR = ("w:shd", *_PPR_AFTER_SHD)
_TCPR_AFTER_SHD = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")


def _shading(fill: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill.lstrip("#"))
    return shading


def _shade_paragraph(paragraph, fill: str) -> None:
    paragraph._p.get_or_add_pPr().insert_element_before(_shading(fill), *_PPR_AFTER_SHD)


def _shade_cell(cell, fill: str) -> None:
    cell._tc.get_or_add_tcPr().insert_element_before(_shading(fill), *_TCPR_AFTER_SHD)


def _border(paragraph, side: str, size: str, space: str, color: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.insert_element_before(borders, *_PPR_AFTER_PBDR)
    edge = OxmlElement(f"w:{side}")
    for key, value in (("w:val", "single"), ("w:sz", size), ("w:space", space), ("w:color", color.lstrip("#"))):
        edge.set(qn(key), value)
    borders.append(edge)


def _left_border(paragraph, color: str) -> None:
    _border(paragraph, "left", "18", "8", color)


def _bottom_border(paragraph, color: str) -> None:
    _border(paragraph, "bottom", "6", "1", color)


def _field(paragraph, instruction: str) -> None:
    """Insert a Word field (e.g. PAGE) that Word updates when rendering."""
    run = paragraph.add_run()
    for tag, text in (("begin", None), (None, instruction), ("separate", None), (None, "1"), ("end", None)):
        if tag:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), tag)
        elif text == instruction:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = f" {instruction} "
        else:
            element = OxmlElement("w:t")
            element.text = text
        run._r.append(element)


def _facts_table(doc, rows: list[tuple[str, str]], theme: Theme) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].width, cells[1].width = Cm(5.5), Cm(10.5)
        label_run = cells[0].paragraphs[0].add_run(label)
        label_run.font.size = Pt(9.5)
        label_run.font.color.rgb = _rgb(theme.muted)
        cells[1].paragraphs[0].add_run(value or "—").font.size = Pt(10.5)
        for cell in cells:
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            _bottom_border(cell.paragraphs[0], theme.line)


def _render_blocks(doc, markdown: str, theme: Theme) -> None:
    for block in parse_blocks(markdown):
        if isinstance(block, Heading):
            level = 2 if block.level <= 2 else 3
            _add_runs(doc.add_paragraph(style=f"Heading {level}"), block.text)
        elif isinstance(block, MdParagraph):
            _add_lines(doc.add_paragraph(), block.lines)
        elif isinstance(block, ListBlock):
            for item in block.items:
                if item.number is not None:
                    paragraph = doc.add_paragraph()
                    paragraph.paragraph_format.left_indent = Cm(0.75 + 0.6 * item.level)
                    paragraph.paragraph_format.first_line_indent = Cm(-0.6)
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.add_run(f"{item.number}.\t")
                    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(0.75 + 0.6 * item.level))
                else:
                    style = "List Bullet" if item.level == 0 else f"List Bullet {min(item.level + 1, 3)}"
                    paragraph = doc.add_paragraph(style=style)
                _add_lines(paragraph, item.lines)
        elif isinstance(block, Quote):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.3)
            _shade_paragraph(paragraph, "F6ECD9")
            _left_border(paragraph, theme.accent)
            _add_lines(paragraph, block.lines)
        elif isinstance(block, MdTable):
            cols = max(1, len(block.header))
            table = doc.add_table(rows=1, cols=cols)
            table.style = "Table Grid"
            for cell, content in zip(table.rows[0].cells, block.header):
                _add_runs(cell.paragraphs[0], content)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                _shade_cell(cell, "EEF3FA")
            for row in block.rows:
                for cell, content in zip(table.add_row().cells, row):
                    _add_runs(cell.paragraphs[0], content)
            doc.add_paragraph()
        elif isinstance(block, Rule):
            _bottom_border(doc.add_paragraph(), theme.line)
        elif isinstance(block, Code):
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(block.text)
            run.font.name = "Consolas"
            run.font.size = Pt(9)


def render_docx(document: ReportDocument, theme: Theme | None = None) -> bytes:
    """Render ``document`` to .docx bytes."""
    theme = theme or Theme()
    doc = Document()
    _setup_styles(doc, theme)
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin, section.bottom_margin = Cm(2.2), Cm(2.0)

    props = doc.core_properties
    props.title = document.title
    props.author = theme.brand_name
    props.subject = f"Human Design — {document.subject.name}"
    props.comments = "Tạo tự động từ dữ liệu đã tính; giờ sinh hiển thị đúng giờ Việt Nam đã khai báo."

    # --- cover ---
    brand = doc.add_paragraph()
    run = brand.add_run(theme.brand_name.upper())
    run.font.size, run.font.color.rgb, run.bold = Pt(10), _rgb(theme.primary), True
    doc.add_paragraph(document.title, style="Title")
    doc.add_paragraph("Thông tin người được phân tích", style="Heading 2")
    _facts_table(doc, subject_lines(document), theme)
    doc.add_paragraph("Bản thiết kế tóm tắt", style="Heading 2")
    _facts_table(doc, chart_facts(document), theme)
    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(18)
    meta_run = meta.add_run(f"Chuẩn bị ngày {prepared_on(document)} · {mode_label(document)}\n"
                            "Giờ sinh hiển thị đúng như đã khai báo (giờ Việt Nam).")
    meta_run.font.size, meta_run.font.color.rgb = Pt(10), _rgb(theme.muted)

    # --- BodyGraph ---
    doc.add_page_break()
    doc.add_paragraph("BodyGraph — Bản đồ năng lượng của bạn", style="Heading 1")
    png = bodygraph_png(document)
    if png:
        doc.add_picture(io.BytesIO(png), height=Cm(20.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = caption.add_run("Trung tâm có màu = năng lượng ổn định, tài năng bẩm sinh · Trung tâm trắng = nơi dễ bị "
                              "ảnh hưởng, cũng là nơi học được trí tuệ · Đen = Personality (ý thức) · Đỏ = Design (vô thức).")
        cap.font.size, cap.font.color.rgb = Pt(9), _rgb(theme.muted)
    else:
        doc.add_paragraph("Không dựng được hình BodyGraph trên máy chủ này — xem file SVG đi kèm báo cáo.")

    # --- content ---
    new_page = breaks_between_sections(document)
    doc.add_page_break()
    for index, report_section in enumerate(included_sections(document)):
        if index and new_page:
            doc.add_page_break()
        doc.add_paragraph(report_section.title, style="Heading 1")
        _render_blocks(doc, report_section.content_markdown, theme)

    # --- header / footer ---
    header = section.header.paragraphs[0]
    header_run = header.add_run(document.title)
    header_run.font.size, header_run.font.color.rgb = Pt(8), _rgb(theme.muted)
    section.different_first_page_header_footer = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run(f"{theme.brand_name} · Trang ")
    footer_run.font.size, footer_run.font.color.rgb = Pt(8), _rgb(theme.muted)
    _field(footer, "PAGE")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


__all__ = ["render_docx"]
